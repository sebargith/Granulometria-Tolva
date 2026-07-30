from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
RUNS_DIR = ROOT / "modelos granulometria" / "clasificador_patches" / "runs"
OUTPUT = ROOT / "docs" / "results" / "Reporte_tests_modelos_granulometria.docx"

CLASSES = ["fino", "medio", "grueso"]


RUN_LABELS = {
    "v001_random_forest": ("Parches 128 iniciales", "Random Forest", "Baseline temprano con pocos datos; resultado exploratorio."),
    "v002_random_forest": ("Parches 128 iniciales", "Random Forest", "Segundo ajuste temprano del baseline; dataset aun pequeno."),
    "ai_raw_etiquetado_ia_frames_50_jobs1": ("IA raw batch001", "Random Forest", "Primer lote etiquetado por IA con patches sin grilla normalizada final."),
    "ai_raw_batch002_jobs1": ("IA raw batch002", "Random Forest", "Segundo lote raw; bajo desempeno en clase fino."),
    "ai_raw_batch001_batch002_jobs1": ("IA raw batch001 + batch002", "Random Forest", "Combinacion de lotes raw con nuevo split."),
    "ai_raw_batch001_batch002_reuse_batch001_split_jobs1": ("IA raw batch001 + batch002", "Random Forest", "Combinacion raw reutilizando split de batch001."),
    "grid96_batch001_ai_prompt_v2_jobs1": ("Grid96 batch001", "Random Forest", "Primera iteracion estable con grilla normalizada 96x96."),
    "resnet18_rgb_grid96_batch001_cuda": ("Grid96 batch001", "ResNet18 RGB", "Primera CNN RGB con transfer learning."),
    "rf_grid96_batch001_batch002_balanced_20260624": ("Grid96 batch001 + batch002", "Random Forest", "Baseline con dos batches y split por frame."),
    "resnet18_rgb_grid96_batch001_batch002_cuda": ("Grid96 batch001 + batch002", "ResNet18 RGB", "CNN entrenada con dos batches."),
    "rf_grid96_batch001_batch002_batch003_reuse_old_split_20260702": ("Grid96 batch001 + batch002 + batch003", "Random Forest", "Prueba combinada, realizada antes de entrenar batch003 por separado."),
    "resnet18_rgb_grid96_batch001_batch002_batch003_reuse_old_split_cuda_20260702": ("Grid96 batch001 + batch002 + batch003", "ResNet18 RGB", "CNN combinada con batch003 agregado."),
    "rf_grid96_batch003_only_20260702": ("Grid96 batch003 original", "Random Forest", "Entrenamiento solo con batch003 antes de revision humana final."),
    "resnet18_rgb_grid96_batch003_only_cuda_20260702": ("Grid96 batch003 original", "ResNet18 RGB", "CNN solo con batch003 antes de revision humana final."),
    "rf_grid96_batch003_revisado_20260702": ("Grid96 batch003 revisado", "Random Forest", "Entrenamiento con batch003 luego de revision manual de etiquetas."),
    "resnet18_rgb_grid96_batch003_revisado_cuda_20260702": ("Grid96 batch003 revisado", "ResNet18 RGB", "Mejor resultado actual; CNN sobre batch003 revisado."),
}


BATCH_ROWS = [
    {
        "nombre": "grid96_batch001_ai_prompt_v2",
        "ruta": r"outputs\patches\grid96_batch001_ai_prompt_v2",
        "descripcion": "Primera version estable: carga normalizada, grilla 96x96 y etiquetado IA con revision.",
        "uso": "1159 patches utiles / 50 frames.",
    },
    {
        "nombre": "grid96_batch002_data2_mixto_ai",
        "ruta": r"outputs\patches\grid96_batch002_data2_mixto_ai",
        "descripcion": "Segundo batch con 90 frames de data/datasets/patch_frames/batch002.",
        "uso": "2305 patches utiles aprox.; combinado con batch001 suma 3464.",
    },
    {
        "nombre": "grid96_batch003_data3_mixto_ai",
        "ruta": r"outputs\patches\grid96_batch003_data3_mixto_ai",
        "descripcion": "Batch003 generado desde data/datasets/patch_frames/batch003 y revisado en la webapp.",
        "uso": "2428 patches utiles / 90 frames. Distribucion: fino 704, medio 903, grueso 821.",
    },
]


INFERENCE_RUNS = [
    (
        "cnn_resnet18_batch003_revisado_prueba_cnn_20260703",
        r"outputs\granulometria_tolva\cnn_resnet18_batch003_revisado_prueba_cnn_20260703\summary.csv",
        "Prueba de 50 frames con el modelo final ResNet18.",
    ),
    (
        "prueba_tecnica_ampliada_cnn_batch003_revisado_20260717_084113",
        r"outputs\prueba_tecnica\prueba_tecnica_ampliada_cnn_batch003_revisado_20260717_084113\summary.csv",
        "Prueba tecnica de 90 frames, 30 por clase visual esperada.",
    ),
]


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None, size: int = 8):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="D9E2EC"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_table(doc, headers, rows, widths=None, font_size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, bold=True, color="0B2545", size=font_size)
        shade_cell(hdr[i], "E8EEF5")
        if widths:
            hdr[i].width = widths[i]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            if widths:
                cells[i].width = widths[i]
    doc.add_paragraph()
    return table


def pct(value):
    if value is None:
        return "-"
    return f"{float(value):.3f}"


def load_runs():
    records = []
    for run in sorted([p for p in RUNS_DIR.iterdir() if p.is_dir()], key=lambda p: p.stat().st_mtime):
        metrics_path = run / "metrics.json"
        summary_path = run / "dataset_summary.json"
        if not metrics_path.exists() or run.name not in RUN_LABELS:
            continue
        metrics = json.loads(metrics_path.read_text(encoding="utf-8"))
        summary = json.loads(summary_path.read_text(encoding="utf-8")) if summary_path.exists() else {}
        input_name, model_name, note = RUN_LABELS[run.name]
        report = metrics.get("test_report", {})
        records.append(
            {
                "run": run.name,
                "input": input_name,
                "model": model_name,
                "patches": summary.get("total_patches", "-"),
                "frames": summary.get("total_frames", "-"),
                "distribution": summary.get("class_distribution", {}),
                "selected": metrics.get("selected_candidate")
                or f"{metrics.get('architecture', 'modelo')} {metrics.get('best_phase', '')} ep. {metrics.get('best_epoch', '')}".strip(),
                "val_f1": metrics.get("best_val_f1_macro")
                or max((x.get("val_f1_macro", 0) for x in metrics.get("candidate_validation", [])), default=None),
                "test_f1": metrics.get("test_f1_macro"),
                "test_acc": metrics.get("test_accuracy"),
                "f1_fino": report.get("fino", {}).get("f1-score"),
                "f1_medio": report.get("medio", {}).get("f1-score"),
                "f1_grueso": report.get("grueso", {}).get("f1-score"),
                "note": note,
                "path": str(run),
            }
        )
    return records


def inference_summary(path: Path):
    if not path.exists():
        return None
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    if not rows:
        return None
    dominant = Counter(row.get("dominant_label", "") for row in rows)
    total_patches = sum(int(float(row.get("patch_count") or 0)) for row in rows)
    def avg(column):
        vals = [float(row.get(column) or 0) for row in rows]
        return sum(vals) / len(vals)
    return {
        "frames": len(rows),
        "patches": total_patches,
        "share_fino": avg("share_fino"),
        "share_medio": avg("share_medio"),
        "share_grueso": avg("share_grueso"),
        "dominant": dict(dominant),
    }


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(text)


def create_doc():
    runs = load_runs()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("Reporte de tests de modelos de granulometria en tolva")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor.from_string("0B2545")
    subtitle = doc.add_paragraph()
    subtitle.add_run(f"Proyecto Granulometria - generado el {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    subtitle.runs[0].font.color.rgb = RGBColor.from_string("555555")

    doc.add_heading("1. Alcance del documento", level=1)
    doc.add_paragraph(
        "Este documento resume los tests realizados para el clasificador de granulometria por patches. "
        "Incluye los inputs o batches usados, los modelos entrenados y las metricas principales disponibles "
        "en los directorios de ejecucion del proyecto."
    )
    add_bullet(doc, "Metrica principal: F1 macro en test, porque pondera de forma equilibrada fino, medio y grueso.")
    add_bullet(doc, "Metrica secundaria: accuracy en test.")
    add_bullet(doc, "Filtro de entrenamiento actual: etiquetas fino/medio/grueso con confianza alta o media.")
    add_bullet(doc, "Split: por frame/source_image para evitar que patches del mismo camion queden repartidos entre train y test.")

    doc.add_heading("2. Inputs y batches utilizados", level=1)
    batch_table_rows = [[row["nombre"], row["ruta"], row["descripcion"], row["uso"]] for row in BATCH_ROWS]
    add_table(
        doc,
        ["Input / batch", "Ruta relativa", "Descripcion", "Uso en tests"],
        batch_table_rows,
        widths=[Inches(1.6), Inches(1.9), Inches(2.3), Inches(1.5)],
        font_size=8,
    )

    doc.add_heading("3. Entrenamientos realizados", level=1)
    doc.add_paragraph(
        "La tabla siguiente consolida todos los entrenamientos con metricas guardadas en "
        "modelos granulometria/clasificador_patches/runs."
    )
    train_rows = []
    for record in runs:
        train_rows.append(
            [
                record["input"],
                record["model"],
                f"{record['patches']} / {record['frames']}",
                record["selected"],
                pct(record["val_f1"]),
                pct(record["test_f1"]),
                pct(record["test_acc"]),
            ]
        )
    add_table(
        doc,
        ["Input", "Modelo", "Patches / frames", "Seleccionado", "Val F1", "Test F1", "Test acc."],
        train_rows,
        widths=[Inches(1.5), Inches(1.0), Inches(0.95), Inches(1.35), Inches(0.55), Inches(0.55), Inches(0.55)],
        font_size=7,
    )

    doc.add_heading("4. Detalle por clase en test", level=1)
    class_rows = []
    for record in runs:
        class_rows.append(
            [
                record["input"],
                record["model"],
                pct(record["f1_fino"]),
                pct(record["f1_medio"]),
                pct(record["f1_grueso"]),
                record["note"],
            ]
        )
    add_table(
        doc,
        ["Input", "Modelo", "F1 fino", "F1 medio", "F1 grueso", "Lectura"],
        class_rows,
        widths=[Inches(1.45), Inches(1.0), Inches(0.55), Inches(0.55), Inches(0.60), Inches(2.75)],
        font_size=7,
    )

    doc.add_heading("5. Comparacion de resultados clave", level=1)
    add_bullet(doc, "La estrategia de grilla normalizada 96x96 supero a los primeros datasets pequenos de patches 128.")
    add_bullet(doc, "Batch001 con Random Forest alcanzo F1 macro test 0.654; la CNN batch001 quedo en 0.574.")
    add_bullet(doc, "Batch001+batch002 no mejoro automaticamente el F1; aumento datos pero tambien variabilidad y ruido de etiquetas.")
    add_bullet(doc, "Batch003 original mostro F1 macro test 0.579 con RF y 0.593 con CNN.")
    add_bullet(doc, "Batch003 revisado fue el salto mas relevante: RF subio a 0.697 y CNN a 0.757 F1 macro test.")

    best = max(runs, key=lambda r: r["test_f1"] or 0)
    doc.add_paragraph(
        f"Resultado actual mas alto: {best['model']} sobre {best['input']} con F1 macro test "
        f"{pct(best['test_f1'])} y accuracy {pct(best['test_acc'])}."
    )

    doc.add_heading("6. Pruebas de inferencia sobre tolva completa", level=1)
    infer_rows = []
    for name, rel_path, note in INFERENCE_RUNS:
        summary = inference_summary(ROOT / rel_path)
        if not summary:
            continue
        dominant_text = ", ".join(f"{k}: {v}" for k, v in sorted(summary["dominant"].items()) if k)
        infer_rows.append(
            [
                name,
                f"{summary['frames']} / {summary['patches']}",
                pct(summary["share_fino"]),
                pct(summary["share_medio"]),
                pct(summary["share_grueso"]),
                dominant_text,
                note,
            ]
        )
    add_table(
        doc,
        ["Run inferencia", "Frames / patches", "% fino prom.", "% medio prom.", "% grueso prom.", "Dominante", "Comentario"],
        infer_rows,
        widths=[Inches(1.45), Inches(0.8), Inches(0.6), Inches(0.65), Inches(0.65), Inches(1.1), Inches(2.1)],
        font_size=7,
    )

    doc.add_heading("7. Conclusiones operativas", level=1)
    add_bullet(doc, "La calidad del etiquetado impacta mas que el tipo de modelo: revisar batch003 produjo la mayor mejora.")
    add_bullet(doc, "Con etiquetas revisadas, CNN ResNet18 supera claramente a Random Forest en F1 macro y accuracy.")
    add_bullet(doc, "Random Forest sigue siendo util como baseline rapido e interpretable, especialmente para validar datasets nuevos.")
    add_bullet(doc, "Las pruebas de inferencia ya generan salidas por frame: mascara, carga normalizada, patches clasificados, overlay y porcentajes fino/medio/grueso.")
    add_bullet(doc, "El siguiente paso recomendable es repetir revision/entrenamiento con mas batches y evaluar sobre un set externo fijo por camion.")

    doc.add_heading("8. Rutas principales", level=1)
    for path in [
        ROOT / "outputs" / "patches",
        ROOT / "modelos granulometria" / "clasificador_patches" / "runs",
        ROOT / "outputs" / "granulometria_tolva",
        ROOT / "tools" / "patch_labeler",
    ]:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(str(path))
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(create_doc())
